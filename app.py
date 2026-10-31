"""MinutesMate — Automated Meeting Minutes Generator (PSID 103).

Flask backend + REST API. The frontend is a proper HTML5/CSS3/JavaScript
single-page web app (templates/index.html + static assets) served by Flask.
Meeting transcripts (pasted or Whisper-transcribed audio) are converted into
an executive summary, key decisions, and action items using GPT or an offline
rule-based engine.
"""
from __future__ import annotations

import os
import tempfile
from io import BytesIO
from pathlib import Path

from dotenv import load_dotenv
from docx import Document
from flask import Flask, jsonify, render_template, request, send_file

from minutes_generator import generate_minutes
from transcribe import transcribe_audio, WhisperUnavailableError

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024  # 200 MB audio uploads


@app.get("/")
def index():
    """Serve the MinutesMate single-page frontend."""
    return render_template("index.html", openai_key_set=bool(os.getenv("OPENAI_API_KEY")))


@app.get("/api/sample")
def sample():
    """Return the bundled sample transcript."""
    path = BASE_DIR / "sample_transcript.txt"
    return jsonify({"transcript": path.read_text(encoding="utf-8") if path.exists() else ""})


@app.post("/api/generate")
def generate():
    """Generate structured minutes from a transcript (GPT or offline mode)."""
    data = request.get_json(silent=True) or {}
    transcript = (data.get("transcript") or "").strip()
    if not transcript:
        return jsonify({"error": "Transcript is empty. Paste a transcript first."}), 400
    api_key = (data.get("api_key") or os.getenv("OPENAI_API_KEY") or "").strip() or None
    model = (data.get("model") or "gpt-4o-mini").strip()
    title = (data.get("title") or "Meeting Minutes").strip()
    try:
        result = generate_minutes(transcript, api_key=api_key, model=model, meeting_title=title)
    except Exception as exc:  # surface a clean error to the frontend
        return jsonify({"error": f"Generation failed: {exc}"}), 500
    return jsonify(result)


@app.post("/api/transcribe")
def transcribe():
    """Transcribe an uploaded meeting recording locally with Whisper."""
    audio = request.files.get("audio")
    if audio is None or audio.filename == "":
        return jsonify({"error": "No audio file received."}), 400
    model_name = request.form.get("model", "base")
    language = (request.form.get("language") or "").strip() or None
    suffix = Path(audio.filename).suffix or ".audio"
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            audio.save(tmp.name)
            temp_path = tmp.name
        text = transcribe_audio(temp_path, model_name=model_name, language=language)
    except WhisperUnavailableError as exc:
        return jsonify({"error": str(exc)}), 503
    except Exception as exc:
        return jsonify({"error": f"Transcription failed: {exc}"}), 500
    finally:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)
    return jsonify({"transcript": text})


@app.post("/api/export/docx")
def export_docx():
    """Build a Word document from generated minutes and return it as a download."""
    data = request.get_json(silent=True) or {}
    title = (data.get("title") or "Meeting Minutes").strip()
    result = data.get("result") or {}

    doc = Document()
    doc.add_heading(f"Meeting Minutes — {title}", 0)
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(result.get("summary", ""))
    doc.add_heading("Key Decisions", level=1)
    for decision in result.get("decisions", []) or ["No explicit decisions detected."]:
        doc.add_paragraph(decision, style="List Bullet")
    doc.add_heading("Action Items", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, ["Task", "Owner", "Due"]):
        cell.text = text
    for item in result.get("action_items", []):
        cells = table.add_row().cells
        cells[0].text = str(item.get("task", ""))
        cells[1].text = str(item.get("owner", "Not specified"))
        cells[2].text = str(item.get("due", "Not specified"))

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name="meeting_minutes.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


if __name__ == "__main__":
    print("MinutesMate running at http://127.0.0.1:5000")
    app.run(host="127.0.0.1", port=5000, debug=False)